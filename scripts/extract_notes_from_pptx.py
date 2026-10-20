from pptx import Presentation
import os

def extract_notes_from_pptx():
    pptx_file = input("Введіть повний шлях до файлу pptx: ").strip()
    
    # Remove quotes if they exist around the path
    if pptx_file.startswith('"') and pptx_file.endswith('"'):
        pptx_file = pptx_file[1:-1]
        
    if not os.path.exists(pptx_file):
        print(f"Помилка: файл '{pptx_file}' не знайдено.")
        return

    # Create notes directory next to the pptx file
    base_dir = os.path.dirname(os.path.abspath(pptx_file))
    notes_dir = os.path.join(base_dir, 'notes')
    os.makedirs(notes_dir, exist_ok=True)
    
    try:
        prs = Presentation(pptx_file)
        
        for i, slide in enumerate(prs.slides, 1):
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                notes_text = notes_slide.notes_text_frame.text
                
                if notes_text.strip():
                    filename = f"note_{i}.txt"
                    filepath = os.path.join(notes_dir, filename)
                    with open(filepath, 'w', encoding='utf-8') as f:
                        f.write(notes_text)
                    print(f"Збережено: {filepath}")
                else:
                    print(f"Слайд {i}: Нотатки відсутні")
            else:
                print(f"Слайд {i}: Нотатки відсутні")
                
        print("\nПроцес видобування завершено.")
    except Exception as e:
        print(f"Сталася помилка під час обробки файлу: {e}")

if __name__ == '__main__':
    extract_notes_from_pptx()
